"""
Excel Generator — 3-Sheet Styled Financial Output
====================================================

Sheet 1: Transactions (exact columns from PDF, styled)
Sheet 2: Summary (opening/closing balance, totals, validation)
Sheet 3: Raw Text (every line from PDF, unprocessed)

Performance: Bulk-write entire DataFrame at once, then style in column sweeps.
Handles 10,000+ rows efficiently.
"""

from __future__ import annotations

import os
import re
from typing import Dict, List, Optional

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter

from utils.logger import get_logger

logger = get_logger("excel_generator")


def _safe_stem(value: str) -> str:
    """ASCII-only, filesystem-safe file stem (drops e.g. U+200E marks)."""
    value = "".join(c for c in value if ord(c) < 128 and c.isprintable())
    return re.sub(r"[^A-Za-z0-9._-]", "_", value).strip("._") or "statement"

# ── Styling constants — Deep Blue premium financial theme ──
_HDR_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
_HDR_FONT = Font(name="Segoe UI", bold=True, color="FFFFFF", size=11)
_DATA_FONT = Font(name="Segoe UI", size=10)
_ALT_FILL = PatternFill(start_color="F0F4FF", end_color="F0F4FF", fill_type="solid")
_SUMMARY_HDR_FILL = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
_SUMMARY_HDR_FONT = Font(name="Segoe UI", bold=True, color="FFFFFF", size=11)
_RAW_HDR_FILL = PatternFill(start_color="4B5563", end_color="4B5563", fill_type="solid")
_BORDER = Border(
    left=Side(style="thin", color="E5E7EB"),
    right=Side(style="thin", color="E5E7EB"),
    top=Side(style="thin", color="E5E7EB"),
    bottom=Side(style="thin", color="E5E7EB"),
)
_CENTER = Alignment(horizontal="center", vertical="center")
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_RIGHT = Alignment(horizontal="right", vertical="center")


def generate_excel(
    dataframes: List[pd.DataFrame],
    file_id: str,
    original_filename: str,
    core_mapping: Optional[Dict[str, str]] = None,
    raw_lines: Optional[List[str]] = None,
    bank_name: str = "UNKNOWN",
    output_dir: Optional[str] = None,
) -> str:
    """Generate a styled 3-sheet Excel output matching the PDF exactly.

    The file is stored as ``<output_dir>/<file_id>.xlsx``; the user-facing
    download name is derived from ``original_filename`` by the download route,
    so user-supplied names never touch the filesystem.
    """
    if output_dir is None:
        from config import settings
        output_dir = settings.output_dir
    os.makedirs(output_dir, exist_ok=True)

    if not core_mapping:
        core_mapping = {}

    out_path = os.path.join(output_dir, f"{_safe_stem(file_id)}.xlsx")

    combined = [df for df in dataframes if not df.empty]
    if combined:
        final = pd.concat(combined, ignore_index=True)
    else:
        final = pd.DataFrame(columns=["No Data"])

    logger.info("Generating 3-sheet Excel: %d rows", len(final))

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        # ═══ Sheet 1: Transactions ═══
        final.to_excel(writer, index=False, sheet_name="Transactions")
        ws_txn = writer.sheets["Transactions"]
        _style_transactions(ws_txn, final, core_mapping)

        # ═══ Sheet 2: Summary ═══
        summary_df = _build_summary(final, core_mapping, bank_name)
        summary_df.to_excel(writer, index=False, sheet_name="Summary")
        ws_sum = writer.sheets["Summary"]
        _style_summary(ws_sum, summary_df)

        # ═══ Sheet 3: Raw Text ═══
        if raw_lines:
            raw_df = pd.DataFrame({"Raw PDF Text": raw_lines})
        else:
            raw_df = pd.DataFrame({"Raw PDF Text": ["(No raw text extracted)"]})
        raw_df.to_excel(writer, index=False, sheet_name="Raw Text")
        ws_raw = writer.sheets["Raw Text"]
        _style_raw(ws_raw, raw_df)

    logger.info("Excel saved: %s (3 sheets, %d txn rows)", os.path.basename(out_path), len(final))
    return out_path


# ═══════════════════════════════════════════════════════════════════════════
# SHEET 1: TRANSACTIONS STYLING
# ═══════════════════════════════════════════════════════════════════════════

def _style_transactions(ws, df: pd.DataFrame, core_mapping: Dict[str, str]) -> None:
    """Apply professional styling to the Transactions sheet."""
    nrows = len(df)
    ncols = len(df.columns)

    date_col = core_mapping.get("Date")
    num_cols = [
        core_mapping.get("Debit"),
        core_mapping.get("Credit"),
        core_mapping.get("Balance"),
    ]
    ref_col = core_mapping.get("Reference")
    desc_col = core_mapping.get("Description")

    # Headers
    for ci in range(1, ncols + 1):
        cell = ws.cell(row=1, column=ci)
        cell.fill = _HDR_FILL
        cell.font = _HDR_FONT
        cell.alignment = _CENTER
        cell.border = _BORDER

    # Data rows — column sweeps
    for ci in range(1, ncols + 1):
        col_name = df.columns[ci - 1]
        is_num = col_name in num_cols
        is_date = col_name == date_col

        for ri in range(2, nrows + 2):
            cell = ws.cell(row=ri, column=ci)
            cell.font = _DATA_FONT
            cell.border = _BORDER

            if is_date or (ref_col and col_name == ref_col):
                cell.alignment = _CENTER
            elif is_num:
                cell.alignment = _RIGHT
            else:
                cell.alignment = _LEFT

            if ri % 2 == 0:
                cell.fill = _ALT_FILL

            if is_num and cell.value is not None and str(cell.value).strip():
                try:
                    cell.value = float(str(cell.value).replace(",", ""))
                    cell.number_format = "#,##0.00"
                except (ValueError, TypeError):
                    pass
            elif is_num:
                cell.value = ""

    # Column widths
    for ci in range(1, ncols + 1):
        letter = get_column_letter(ci)
        col_name = df.columns[ci - 1]

        if col_name == date_col:
            width = 14
        elif col_name in num_cols:
            width = 18
        elif col_name == ref_col:
            width = 18
        elif col_name == desc_col:
            width = 55
        else:
            width = 22

        ws.column_dimensions[letter].width = width

    ws.freeze_panes = "A2"


# ═══════════════════════════════════════════════════════════════════════════
# SHEET 2: SUMMARY
# ═══════════════════════════════════════════════════════════════════════════

def _build_summary(df: pd.DataFrame, core_mapping: Dict[str, str], bank_name: str) -> pd.DataFrame:
    """Build a summary DataFrame with key financial stats."""
    rows: List[Dict[str, str]] = []

    rows.append({"Metric": "Bank", "Value": bank_name})
    rows.append({"Metric": "Total Transactions", "Value": str(len(df))})

    debit_col = core_mapping.get("Debit")
    credit_col = core_mapping.get("Credit")
    balance_col = core_mapping.get("Balance")

    if debit_col and debit_col in df.columns:
        total_debit = _safe_sum(df[debit_col])
        rows.append({"Metric": "Total Debits", "Value": f"{total_debit:,.2f}"})

    if credit_col and credit_col in df.columns:
        total_credit = _safe_sum(df[credit_col])
        rows.append({"Metric": "Total Credits", "Value": f"{total_credit:,.2f}"})

    if balance_col and balance_col in df.columns:
        balances = df[balance_col].dropna()
        balances_num = pd.to_numeric(
            balances.astype(str).str.replace(",", "", regex=False),
            errors="coerce",
        ).dropna()
        if len(balances_num) > 0:
            rows.append({"Metric": "Opening Balance", "Value": f"{balances_num.iloc[0]:,.2f}"})
            rows.append({"Metric": "Closing Balance", "Value": f"{balances_num.iloc[-1]:,.2f}"})

    date_col = core_mapping.get("Date")
    if date_col and date_col in df.columns:
        dates = df[date_col].dropna()
        if len(dates) > 0:
            rows.append({"Metric": "First Date", "Value": str(dates.iloc[0])})
            rows.append({"Metric": "Last Date", "Value": str(dates.iloc[-1])})

    # Validation check
    if debit_col and credit_col and balance_col:
        errors = _check_balance_continuity(df, core_mapping)
        rows.append({"Metric": "Balance Validation Errors", "Value": str(errors)})

    return pd.DataFrame(rows)


def _safe_sum(series: pd.Series) -> float:
    """Sum a series of mixed string/float values safely."""
    total = 0.0
    for val in series:
        if val is None or str(val).strip() == "":
            continue
        try:
            total += abs(float(str(val).replace(",", "")))
        except (ValueError, TypeError):
            pass
    return total


def _check_balance_continuity(df: pd.DataFrame, core_mapping: Dict[str, str]) -> int:
    """Count balance continuity violations (±1 tolerance)."""
    debit_col = core_mapping.get("Debit", "")
    credit_col = core_mapping.get("Credit", "")
    balance_col = core_mapping.get("Balance", "")
    errors = 0

    prev_balance = None
    for _, row in df.iterrows():
        try:
            bal = float(str(row.get(balance_col, "")).replace(",", ""))
        except (ValueError, TypeError):
            continue

        if prev_balance is not None:
            try:
                debit = float(str(row.get(debit_col, "0")).replace(",", "") or "0")
            except (ValueError, TypeError):
                debit = 0.0
            try:
                credit = float(str(row.get(credit_col, "0")).replace(",", "") or "0")
            except (ValueError, TypeError):
                credit = 0.0

            expected = prev_balance + credit - debit
            if abs(expected - bal) > 1.0:
                errors += 1

        prev_balance = bal

    return errors


# ═══════════════════════════════════════════════════════════════════════════
# STYLING HELPERS
# ═══════════════════════════════════════════════════════════════════════════

def _style_summary(ws, df: pd.DataFrame) -> None:
    """Style the Summary sheet."""
    for ci in range(1, 3):
        cell = ws.cell(row=1, column=ci)
        cell.fill = _SUMMARY_HDR_FILL
        cell.font = _SUMMARY_HDR_FONT
        cell.alignment = _CENTER
        cell.border = _BORDER

    for ri in range(2, len(df) + 2):
        for ci in range(1, 3):
            cell = ws.cell(row=ri, column=ci)
            cell.font = _DATA_FONT
            cell.border = _BORDER
            cell.alignment = _LEFT if ci == 1 else _RIGHT
            if ri % 2 == 0:
                cell.fill = _ALT_FILL

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 25
    ws.freeze_panes = "A2"


def _style_raw(ws, df: pd.DataFrame) -> None:
    """Style the Raw Text sheet."""
    cell = ws.cell(row=1, column=1)
    cell.fill = _RAW_HDR_FILL
    cell.font = _HDR_FONT
    cell.alignment = _CENTER
    cell.border = _BORDER

    for ri in range(2, len(df) + 2):
        cell = ws.cell(row=ri, column=1)
        cell.font = Font(name="Consolas", size=9)
        cell.alignment = _LEFT
        cell.border = _BORDER
        if ri % 2 == 0:
            cell.fill = _ALT_FILL

    ws.column_dimensions["A"].width = 120
    ws.freeze_panes = "A2"
